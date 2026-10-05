from io import BytesIO

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from fpdf import FPDF


def format_docx(
    content,
    document_type=None,
    brand_name=None,
    logo_bytes=None
):
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)

    normal_style._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Times New Roman"
    )

    if logo_bytes:
        try:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(
                BytesIO(logo_bytes),
                width=Inches(1.2)
            )
        except Exception:
            pass

    if brand_name:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(str(brand_name))
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)

    if document_type:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(str(document_type).upper())
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(16)

    for line in content.splitlines():
        line = line.strip()

        if not line:
            document.add_paragraph()
            continue

        if (
            line.isupper()
            or line.startswith("1.")
            or line.startswith("2.")
            or line.startswith("3.")
            or line.startswith("4.")
            or line.startswith("5.")
            or line.startswith("6.")
            or line.startswith("7.")
            or line.startswith("8.")
            or line.startswith("9.")
        ):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

        elif line.startswith("-") or line.startswith("•"):
            paragraph = document.add_paragraph(
                style="List Bullet"
            )
            run = paragraph.add_run(
                line.lstrip("-• ").strip()
            )
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

        else:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer_run = footer_paragraph.add_run(
        "Generated using LegalEase - AI-Assisted Legal Document Generator"
    )
    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(8)

    output = BytesIO()
    document.save(output)
    output.seek(0)

    return output.getvalue()


def format_pdf(
    content,
    document_type=None,
    brand_name=None,
    logo_bytes=None
):
    pdf = FPDF()
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    usable_width = pdf.w - pdf.l_margin - pdf.r_margin

    if brand_name:
        pdf.set_font("Times", "B", 14)
        pdf.multi_cell(
            usable_width,
            8,
            str(brand_name),
            align="C"
        )
        pdf.ln(3)

    if document_type:
        pdf.set_font("Times", "B", 16)
        pdf.multi_cell(
            usable_width,
            9,
            str(document_type).upper(),
            align="C"
        )
        pdf.ln(5)

    pdf.set_font("Times", "", 11)

    for line in content.splitlines():
        line = line.strip()

        if not line:
            pdf.ln(5)
            continue

        words = line.split()
        safe_line = ""

        for word in words:
            if pdf.get_string_width(word) > usable_width:
                if safe_line:
                    pdf.multi_cell(
                        usable_width,
                        7,
                        safe_line
                    )
                    safe_line = ""

                chunk = ""

                for char in word:
                    if pdf.get_string_width(
                        chunk + char
                    ) > usable_width:
                        pdf.multi_cell(
                            usable_width,
                            7,
                            chunk
                        )
                        chunk = char
                    else:
                        chunk += char

                if chunk:
                    safe_line = chunk

            else:
                test_line = (
                    safe_line + " " + word
                    if safe_line
                    else word
                )

                if pdf.get_string_width(
                    test_line
                ) > usable_width:
                    if safe_line:
                        pdf.multi_cell(
                            usable_width,
                            7,
                            safe_line
                        )
                    safe_line = word
                else:
                    safe_line = test_line

        if safe_line:
            if (
                line.isupper()
                or line.startswith("1.")
                or line.startswith("2.")
                or line.startswith("3.")
                or line.startswith("4.")
                or line.startswith("5.")
                or line.startswith("6.")
                or line.startswith("7.")
                or line.startswith("8.")
                or line.startswith("9.")
            ):
                pdf.set_font("Times", "B", 12)
                pdf.multi_cell(
                    usable_width,
                    8,
                    safe_line
                )
                pdf.set_font("Times", "", 11)
            else:
                pdf.multi_cell(
                    usable_width,
                    7,
                    safe_line
                )

    pdf.set_y(-15)
    pdf.set_font("Times", "", 8)
    pdf.cell(
        usable_width,
        10,
        "Generated using LegalEase - AI-Assisted Legal Document Generator",
        align="C"
    )

    return bytes(pdf.output())